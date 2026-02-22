"""
SonIA Core ÃÂ¢ÃÂÃÂ Warehouse Excel Parser
Parses warehouse Excel files to extract data per dropshipper:
  - Unique orders, box types/counts, SKUs sold, tracking numbers.

Excel structure:
  - Tab "PICKING LIST": global product picking summary
  - Tab "CAJAS": box type counts (global)
  - Tab "[BRAND] - PACK LIST": per-brand order detail with boxes, products, AWBs
  - Tab "PRODUCTS SOLD": SKU quantities grouped by brand
"""

import logging
import re
from typing import Dict, List, Any, Optional
from pathlib import Path

import openpyxl

logger = logging.getLogger(__name__)


class WarehouseParser:
    """Parse a warehouse Excel file and extract structured data per dropshipper."""

    def __init__(self, file_path: str):
        self.file_path = file_path
        self.wb = None
        self.sheet_names: List[str] = []

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def parse(self) -> Dict[str, Dict[str, Any]]:
        """
        Parse the warehouse Excel and return data grouped by dropshipper.

        Returns:
            {
                "Dios Mio Coffee": {
                    "unique_orders": 8,
                    "boxes": {"CAJA PEQUEÃÂÃÂA": 5, "CAJA MEDIANA": 2},
                    "total_boxes": 7,
                    "skus": {"DMC12BMGS": 4, "DMC12BLGS": 2, ...},
                    "total_skus_sold": 12,
                    "tracking_numbers": ["888902157098", ...],
                },
                ...
            }
        """
        self.wb = openpyxl.load_workbook(self.file_path, data_only=True)
        self.sheet_names = self.wb.sheetnames
        logger.info(f"Opened warehouse file: {self.file_path}")
        logger.info(f"Sheets found: {self.sheet_names}")

        # 1. Identify brands from PACK LIST tabs
        brands = self._identify_brands()
        logger.info(f"Brands identified: {brands}")

        if not brands:
            # Fallback: try Hoja1 format (COOCENTRAL, DON MAIZ B2B, etc.)
            if "Hoja1" in self.sheet_names:
                logger.info("No PACK LIST tabs found, trying Hoja1 format")
                return self._parse_hoja1()
            raise ValueError("No brand PACK LIST tabs found in warehouse file")

        result = {}

        for brand in brands:
            pack_list_sheet = f"{brand} - PACK LIST"
            brand_data = {
                "unique_orders": 0,
                "boxes": {},
                "total_boxes": 0,
                "skus": {},
                "total_skus_sold": 0,
                "tracking_numbers": [],
                "dispatch_date": None,
            }

            # 2. Parse the PACK LIST for this brand (orders, boxes, AWBs)
            if pack_list_sheet in self.sheet_names:
                pack_data = self._parse_pack_list(pack_list_sheet)
                brand_data["unique_orders"] = pack_data["unique_orders"]
                brand_data["boxes"] = pack_data["boxes"]
                brand_data["total_boxes"] = sum(pack_data["boxes"].values())
                brand_data["tracking_numbers"] = pack_data["tracking_numbers"]
                brand_data["dispatch_date"] = pack_data.get("dispatch_date")

            # 3. Parse PRODUCTS SOLD for this brand (SKUs)
            if "PRODUCTS SOLD" in self.sheet_names:
                skus = self._parse_products_sold_for_brand(brand)
                brand_data["skus"] = skus
                brand_data["total_skus_sold"] = sum(skus.values())
            elif "PICKING LIST" in self.sheet_names:
                skus = self._parse_picking_list()
                brand_data["skus"] = skus
                brand_data["total_skus_sold"] = sum(skus.values())

            result[brand] = brand_data
            logger.info(
                f"  {brand}: {brand_data['unique_orders']} orders, "
                f"{brand_data['total_boxes']} boxes, "
                f"{brand_data['total_skus_sold']} SKUs"
            )

        self.wb.close()
        return result

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------

    def _identify_brands(self) -> List[str]:
        """Find brands from sheet names matching '[BRAND] - PACK LIST'."""
        brands = []
        for name in self.sheet_names:
            match = re.match(r"^(.+?)\s*-\s*PACK LIST$", name, re.IGNORECASE)
            if match:
                brands.append(match.group(1).strip())
        return brands

    def _parse_pack_list(self, sheet_name: str) -> Dict[str, Any]:
        """
        Parse a brand's PACK LIST tab.

        Structure:
            Row with ORDEN # starts a new order block.
            Following rows (empty ORDEN #) are additional items in same order.
            Row with 'TOTAL CAJAS:' in col C ends the order block.
            Col D = TIPO CAJA, Col I = AWB (tracking number).
        """
        ws = self.wb[sheet_name]
        orders = set()
        boxes = {}
        tracking_numbers = []

        # Find the header row
        header_row = None
        dispatch_date = None
        # Extract FECHA CREACION from header rows
        for r in ws.iter_rows(min_row=1, max_row=8, values_only=True):
            str_vals = [str(v or "").strip() for v in r]
            if str_vals and "FECHA" in str_vals[0].upper() and "CREACI" in str_vals[0].upper():
                raw_val = r[1] if len(r) > 1 else None
                logger.info(f"Found FECHA row in {sheet_name}, raw value: {raw_val} (type: {type(raw_val).__name__})")
                if raw_val is not None:
                    try:
                        from datetime import datetime as _dt, date as _date
                        # Handle Excel datetime objects directly
                        if isinstance(raw_val, _dt):
                            dispatch_date = raw_val.date()
                        elif isinstance(raw_val, _date):
                            dispatch_date = raw_val
                        else:
                            raw_date = str(raw_val).strip()
                            for fmt in ("%d %b %Y", "%d %B %Y", "%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
                                try:
                                    dispatch_date = _dt.strptime(raw_date, fmt).date()
                                    break
                                except ValueError:
                                    continue
                        if dispatch_date is None:
                            logger.warning(f"Could not parse date from: {raw_val} (type: {type(raw_val).__name__})")
                        else:
                            logger.info(f"Extracted dispatch_date: {dispatch_date} from {sheet_name}")
                    except Exception as e:
                        logger.warning(f"Error parsing date from {sheet_name}: {e}")
                break

        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=False), 1):
            vals = [c.value for c in row]
            if vals and str(vals[0] or "").strip().upper() == "ORDEN #":
                header_row = i
                break

        if header_row is None:
            logger.warning(f"No header row found in {sheet_name}")
            return {"unique_orders": 0, "boxes": {}, "tracking_numbers": []}

        current_order = None

        for row in ws.iter_rows(min_row=header_row + 1, values_only=False):
            vals = [c.value for c in row]

            # Skip completely empty rows
            if not any(v is not None and str(v).strip() for v in vals):
                continue

            col_a = str(vals[0] or "").strip()   # ORDEN #
            col_c = str(vals[2] or "").strip()   # May have TOTAL CAJAS:
            col_d = str(vals[3] or "").strip()   # TIPO CAJA or count
            col_i = str(vals[8] or "").strip() if len(vals) > 8 else ""  # AWB

            # New order starts when col A has an order number
            if col_a and col_a.startswith("#"):
                current_order = col_a
                orders.add(current_order)

                # This row has a box type and possibly AWB
                if col_d and col_d.upper() not in ("", "TIPO CAJA"):
                    box_type = col_d.strip()
                    boxes[box_type] = boxes.get(box_type, 0) + 1

                if col_i and col_i not in ("", "AWB"):
                    tracking_numbers.append(col_i)

            elif "TOTAL CAJAS:" in col_c.upper():
                # End of order block ÃÂ¢ÃÂÃÂ skip
                continue

            else:
                # Continuation row (additional boxes/items for current order)
                if col_d and col_d.upper() not in ("", "TIPO CAJA"):
                    # Check if this is a box type (not a number)
                    try:
                        float(col_d)
                        # It's a number (total count), skip
                    except ValueError:
                        box_type = col_d.strip()
                        boxes[box_type] = boxes.get(box_type, 0) + 1

                if col_i and col_i not in ("", "AWB"):
                    tracking_numbers.append(col_i)

        return {
            "unique_orders": len(orders),
            "boxes": boxes,
            "tracking_numbers": tracking_numbers,
            "dispatch_date": dispatch_date,
        }


    def _parse_hoja1(self) -> Dict[str, Dict[str, Any]]:
        """
        Parse Hoja1-only format (COOCENTRAL, DON MAIZ B2B, etc.).

        Structure:
            Row 4: FECHA CREACION (may be empty)
            Row 7: Header row (ORDEN #, Canal, DESTINATARIO, ZIPCODE, DIRECCION, CIUDAD, TELEFONO, TIPO CAJA, NOMBRE PRODUCTO, CANTIDAD UNIDADES)
            Data rows: order blocks similar to PACK LIST
            Rows with 'TOTAL CAJA:' in col D end order blocks.

        Brand is inferred from the filename (folder name passed by the uploader).
        Since we don't know the brand from Hoja1, we use a generic brand name
        that the processor will try to resolve from the upload context.
        """
        ws = self.wb["Hoja1"]
        orders = set()
        boxes = {}
        tracking_numbers = []
        skus = {}
        dispatch_date = None

        # Extract FECHA CREACION from header rows
        for r in ws.iter_rows(min_row=1, max_row=8, values_only=True):
            str_vals = [str(v or "").strip() for v in r]
            if str_vals and "FECHA" in str_vals[0].upper() and "CREACI" in str_vals[0].upper():
                raw_val = r[1] if len(r) > 1 else None
                if raw_val is not None:
                    try:
                        from datetime import datetime as _dt, date as _date
                        if isinstance(raw_val, _dt):
                            dispatch_date = raw_val.date()
                        elif isinstance(raw_val, _date):
                            dispatch_date = raw_val
                        else:
                            raw_date = str(raw_val).strip()
                            for fmt in ("%d %b %Y", "%d %B %Y", "%Y-%m-%d", "%d/%m/%Y"):
                                try:
                                    dispatch_date = _dt.strptime(raw_date, fmt).date()
                                    break
                                except ValueError:
                                    continue
                    except Exception as e:
                        logger.warning(f"Error parsing Hoja1 date: {e}")
                break

        # Find header row
        header_row = None
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), 1):
            col_a = str(row[0] or "").strip().upper() if row[0] else ""
            if "ORDEN" in col_a:
                header_row = i
                break

        if not header_row:
            logger.warning("No header row found in Hoja1")
            return {}

        # Parse data rows
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            col_a = str(row[0] or "").strip() if row[0] else ""  # ORDEN #
            col_d = str(row[3] or "").strip() if len(row) > 3 and row[3] else ""  # ZIPCODE or TOTAL CAJA
            col_h = str(row[7] or "").strip() if len(row) > 7 and row[7] else ""  # TIPO CAJA
            col_i = str(row[8] or "").strip() if len(row) > 8 and row[8] else ""  # NOMBRE PRODUCTO
            col_j = row[9] if len(row) > 9 else None  # CANTIDAD UNIDADES

            if "TOTAL CAJA" in col_d.upper():
                continue

            if col_a and col_a.startswith("#"):
                orders.add(col_a)

            if col_h and col_h.upper() not in ("", "TIPO CAJA"):
                boxes[col_h] = boxes.get(col_h, 0) + 1

            # Extract product as SKU (use product name since no SAP code in Hoja1)
            if col_i and col_i.upper() not in ("", "NOMBRE PRODUCTO") and col_j is not None:
                try:
                    qty = int(col_j)
                    skus[col_i] = skus.get(col_i, 0) + qty
                except (ValueError, TypeError):
                    pass

        # Use "UNKNOWN" as brand - processor will resolve from upload context
        brand_name = "UNKNOWN"
        logger.info(f"  Hoja1 parsed: {len(orders)} orders, {sum(boxes.values())} boxes, {sum(skus.values())} SKUs")

        return {
            brand_name: {
                "unique_orders": len(orders),
                "boxes": boxes,
                "total_boxes": sum(boxes.values()),
                "skus": skus,
                "total_skus_sold": sum(skus.values()),
                "tracking_numbers": tracking_numbers,
                "dispatch_date": dispatch_date,
            }
        }

    def _parse_picking_list(self) -> Dict[str, int]:
        """
        Fallback: parse PICKING LIST tab to extract SKUs when PRODUCTS SOLD is missing.

        Structure:
            Header row has 'SAP TERC' in col A.
            Data rows: SAP code in col A, quantity in col D.
            Ends at row with 'TOTAL' or empty SAP.
        """
        ws = self.wb["PICKING LIST"]
        skus = {}
        in_data = False

        for row in ws.iter_rows(min_row=1, values_only=True):
            col_a = str(row[0] or "").strip() if row[0] is not None else ""
            col_d = row[3] if len(row) > 3 else None

            if not in_data:
                if "SAP" in col_a.upper():
                    in_data = True
                    continue
            else:
                if not col_a or "TOTAL" in col_a.upper():
                    break
                if col_d is not None:
                    try:
                        qty = int(col_d)
                        skus[col_a] = qty
                    except (ValueError, TypeError):
                        continue

        logger.info(f"  Parsed PICKING LIST: {len(skus)} SKUs, total {sum(skus.values())} units")
        return skus

    def _parse_products_sold_for_brand(self, brand: str) -> Dict[str, int]:
        """
        Parse the PRODUCTS SOLD tab for a specific brand.

        Structure:
            Brand name appears alone in col A as a section header.
            Next row is column headers: SKU, NOMBRE PRODUCTO, TIPO, CANTIDAD
            Then data rows until a row with 'TOTAL:' in col C.
        """
        ws = self.wb["PRODUCTS SOLD"]
        skus = {}
        in_brand_section = False
        expect_header = False

        for row in ws.iter_rows(min_row=1, values_only=False):
            vals = [c.value for c in row]

            if not any(v is not None and str(v).strip() for v in vals):
                continue

            col_a = str(vals[0] or "").strip()
            col_c = str(vals[2] or "").strip() if len(vals) > 2 else ""
            col_d = vals[3] if len(vals) > 3 else None

            if expect_header:
                # This should be the header row (SKU, NOMBRE, TIPO, CANTIDAD)
                if col_a.upper() == "SKU":
                    expect_header = False
                    in_brand_section = True
                    continue
                else:
                    expect_header = False
                    continue

            if in_brand_section:
                # Check for end of section (TOTAL row)
                if "TOTAL" in col_c.upper():
                    in_brand_section = False
                    continue

                # Data row: SKU in col A, quantity in col D
                if col_a and col_d is not None:
                    try:
                        qty = int(col_d)
                        skus[col_a] = qty
                    except (ValueError, TypeError):
                        pass
                continue

            # Look for brand section header
            # Brand name matching: case-insensitive, partial match
            if col_a and self._brand_matches(col_a, brand):
                # Check that cols B, C, D are empty (brand header is alone)
                other_vals = [str(v or "").strip() for v in vals[1:4] if v is not None]
                if not any(other_vals):
                    expect_header = True

        return skus

    def _brand_matches(self, cell_value: str, brand: str) -> bool:
        """Check if a cell value matches a brand name (flexible matching)."""
        cell_lower = cell_value.lower().strip()
        brand_lower = brand.lower().strip()

        # Exact match
        if cell_lower == brand_lower:
            return True

        # Brand is contained in cell or vice versa
        if brand_lower in cell_lower or cell_lower in brand_lower:
            return True

        return False
